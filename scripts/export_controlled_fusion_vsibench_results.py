#!/usr/bin/env python3
"""Export the official controlled-fusion VSI-Bench campaign results.

The Excel workbook contains a human-readable summary sheet and a detailed
per-question-type sheet.  Matching CSV files are emitted because CSV itself
cannot contain multiple worksheets.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ARCHITECTURES = {
    "A": "VLM3R pre-LLM Cross-attn",
    "B": "pre-projector Add, dec12",
    "C": "L0 Cross-attn, dec12",
    "D": "L0 Add, dec12",
    "E": "L0/1/2 Add, dec12x3, site-specific",
    "F": "L0/1/2 Add, dec6/9/12",
    "G": "L0/1/2 Cross-attn, dec6/9/12",
    "H": "L0/1/2 Cross-attn, dec12x3",
    "Reproduction_2": "Independent trained-checkpoint diagnostic",
}

CONTROLLED_IDS = ("B", "C", "D", "E", "H")
REFERENCE_IDS = ("A", "F", "G")
SUMMARY_IDS = ("A", "B", "C", "D", "E", "F", "G", "H", "Reproduction_2")
MILESTONE_STAGES = ("p01", "p05", "p25", "p50")
CONTROLLED_STAGES = MILESTONE_STAGES + ("final",)
STAGE_LABELS = {
    "p01": "1%",
    "p05": "5%",
    "p25": "25%",
    "p50": "50%",
    "final": "Final",
}
STAGE_STEPS = {"p01": 17, "p05": 82, "p25": 406, "p50": 811, "final": 1622}
STAGE_RATIOS = {"p01": 0.01, "p05": 0.05, "p25": 0.25, "p50": 0.50, "final": 1.0}

SUBTASKS = (
    "room_size_estimation",
    "object_size_estimation",
    "object_counting",
    "object_abs_distance",
    "object_rel_distance",
    "object_rel_direction_easy",
    "object_rel_direction_medium",
    "object_rel_direction_hard",
    "route_planning",
    "obj_appearance_order",
)
DIRECTION_SUBTASKS = (
    "object_rel_direction_easy",
    "object_rel_direction_medium",
    "object_rel_direction_hard",
)
NA_SUBTASKS = {
    "room_size_estimation",
    "object_size_estimation",
    "object_counting",
    "object_abs_distance",
}
SUBTASK_LABELS = {
    "room_size_estimation": "Room Size",
    "object_size_estimation": "Object Size",
    "object_counting": "Object Counting",
    "object_abs_distance": "Absolute Distance",
    "object_rel_distance": "Relative Distance",
    "object_rel_direction_easy": "Relative Direction Easy",
    "object_rel_direction_medium": "Relative Direction Medium",
    "object_rel_direction_hard": "Relative Direction Hard",
    "route_planning": "Route Planning",
    "obj_appearance_order": "Appearance Order",
}


@dataclass(frozen=True)
class RunSpec:
    campaign: str
    architecture_id: str
    stage: str
    directory: Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--eval-root",
        type=Path,
        required=True,
        help="Root containing controlled_fusion_official and reference_fusion_half_official",
    )
    parser.add_argument(
        "--reproduction-dir",
        type=Path,
        help="Optional full Reproduction_2 evaluation directory",
    )
    parser.add_argument("--output-prefix", type=Path, required=True)
    return parser.parse_args()


def _files_under(directory: Path, filename: str) -> list[Path]:
    return sorted(directory.rglob(filename), key=lambda path: path.stat().st_mtime)


def _has_complete_result(directory: Path) -> bool:
    return bool(_files_under(directory, "results.json") and _files_under(directory, "vsibench.json"))


def _unique_campaign_root(parent: Path, pattern: str) -> Path:
    matches = sorted(path for path in parent.glob(pattern) if path.is_dir())
    complete = [path for path in matches if _files_under(path, "results.json")]
    if len(complete) != 1:
        raise RuntimeError(
            f"Expected one completed campaign root for {parent / pattern}, found: {complete}"
        )
    return complete[0]


def _stage_directory(campaign_root: Path, stage: str) -> Path:
    patterns = (f"{stage}_*",) if stage != "final" else ("final_*", "eval_*")
    candidates: list[Path] = []
    for pattern in patterns:
        candidates.extend(path for path in campaign_root.glob(pattern) if path.is_dir())
    completed = [path for path in candidates if _has_complete_result(path)]
    if not completed:
        raise RuntimeError(f"No completed {stage} result under {campaign_root}")
    return max(
        completed,
        key=lambda path: _files_under(path, "results.json")[-1].stat().st_mtime,
    )


def discover_runs(eval_root: Path, reproduction_dir: Path | None) -> list[RunSpec]:
    controlled_parent = eval_root / "controlled_fusion_official"
    reference_parent = eval_root / "reference_fusion_half_official"
    runs: list[RunSpec] = []

    for architecture_id in CONTROLLED_IDS:
        root = _unique_campaign_root(controlled_parent, f"controlled_{architecture_id}_*")
        for stage in CONTROLLED_STAGES:
            runs.append(RunSpec("controlled", architecture_id, stage, _stage_directory(root, stage)))

    for architecture_id in REFERENCE_IDS:
        root = _unique_campaign_root(reference_parent, f"reference_{architecture_id}_*")
        for stage in MILESTONE_STAGES:
            runs.append(RunSpec("reference_half", architecture_id, stage, _stage_directory(root, stage)))

    if reproduction_dir is not None:
        if not _has_complete_result(reproduction_dir):
            raise RuntimeError(f"Incomplete reproduction result: {reproduction_dir}")
        runs.append(RunSpec("reproduction", "Reproduction_2", "final", reproduction_dir))
    return runs


def _score_record(log: dict[str, Any]) -> dict[str, Any]:
    value = log.get("vsibench_score")
    if isinstance(value, dict):
        return value
    value = log.get("doc")
    return value if isinstance(value, dict) else {}


def _subtask_score(record: dict[str, Any], question_type: str) -> float:
    key = "MRA:.5:.95:.05" if question_type in NA_SUBTASKS else "accuracy"
    value = record.get(key)
    if not isinstance(value, (int, float)):
        raise RuntimeError(f"Missing numeric {key} for {question_type}: {record}")
    return float(value)


def load_run(spec: RunSpec) -> dict[str, Any]:
    results_paths = _files_under(spec.directory, "results.json")
    detail_paths = _files_under(spec.directory, "vsibench.json")
    if len(results_paths) != 1 or len(detail_paths) != 1:
        raise RuntimeError(
            f"Expected one results.json and vsibench.json under {spec.directory}; "
            f"found {results_paths} and {detail_paths}"
        )
    results_path = results_paths[0]
    detail_path = detail_paths[0]
    results = json.loads(results_path.read_text(encoding="utf-8"))
    detail = json.loads(detail_path.read_text(encoding="utf-8"))

    official_score = float(results["results"]["vsibench"]["vsibench_score,none"])
    logs = detail.get("logs")
    if not isinstance(logs, list) or len(logs) != 5130:
        raise RuntimeError(f"Expected 5130 detail records in {detail_path}, found {len(logs or [])}")

    values: dict[str, list[float]] = defaultdict(list)
    for log in logs:
        record = _score_record(log)
        question_type = record.get("question_type")
        if question_type not in SUBTASKS:
            raise RuntimeError(f"Unexpected VSI-Bench question type: {question_type!r}")
        values[question_type].append(_subtask_score(record, question_type))

    missing = [question_type for question_type in SUBTASKS if not values[question_type]]
    if missing:
        raise RuntimeError(f"Missing subtasks in {detail_path}: {missing}")
    per_subtask = {
        question_type: sum(values[question_type]) / len(values[question_type]) * 100.0
        for question_type in SUBTASKS
    }
    direction_grouped = sum(per_subtask[name] for name in DIRECTION_SUBTASKS) / 3.0
    official_components = [
        value for name, value in per_subtask.items() if name not in DIRECTION_SUBTASKS
    ] + [direction_grouped]
    reconstructed = sum(official_components) / len(official_components)
    if not math.isclose(official_score, reconstructed, rel_tol=0.0, abs_tol=1e-9):
        raise RuntimeError(
            f"Official score mismatch for {spec.architecture_id}/{spec.stage}: "
            f"results={official_score}, reconstructed={reconstructed}"
        )

    return {
        "Campaign": spec.campaign,
        "ID": spec.architecture_id,
        "Architecture": ARCHITECTURES[spec.architecture_id],
        "Stage": STAGE_LABELS[spec.stage],
        "Stage Key": spec.stage,
        "Training Fraction": "" if spec.campaign == "reproduction" else STAGE_RATIOS[spec.stage],
        "Global Step": "" if spec.campaign == "reproduction" else STAGE_STEPS[spec.stage],
        "Overall": official_score,
        "Direction Grouped": direction_grouped,
        "Total Samples": len(logs),
        "Evaluation Directory": str(spec.directory),
        "Results JSON": str(results_path),
        "Detail JSON": str(detail_path),
        **{SUBTASK_LABELS[name]: per_subtask[name] for name in SUBTASKS},
        **{f"{SUBTASK_LABELS[name]} Samples": len(values[name]) for name in SUBTASKS},
    }


def summary_rows(detail_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    lookup = {(row["ID"], row["Stage Key"]): row["Overall"] for row in detail_rows}
    rows = []
    for architecture_id in SUMMARY_IDS:
        campaign = (
            "controlled"
            if architecture_id in CONTROLLED_IDS
            else "reference_half"
            if architecture_id in REFERENCE_IDS
            else "reproduction"
        )
        rows.append(
            {
                "Campaign": campaign,
                "ID": architecture_id,
                "Architecture": ARCHITECTURES[architecture_id],
                "1%": lookup.get((architecture_id, "p01"), ""),
                "5%": lookup.get((architecture_id, "p05"), ""),
                "25%": lookup.get((architecture_id, "p25"), ""),
                "50%": lookup.get((architecture_id, "p50"), ""),
                "100% / Final": lookup.get((architecture_id, "final"), ""),
                "Notes": (
                    "Trained to 50% by design"
                    if architecture_id in REFERENCE_IDS
                    else "Independent diagnostic; no milestone checkpoints"
                    if architecture_id == "Reproduction_2"
                    else ""
                ),
            }
        )
    return rows


SUMMARY_HEADERS = (
    "Campaign",
    "ID",
    "Architecture",
    "1%",
    "5%",
    "25%",
    "50%",
    "100% / Final",
    "Notes",
)
DETAIL_HEADERS = (
    "Campaign",
    "ID",
    "Architecture",
    "Stage",
    "Training Fraction",
    "Global Step",
    "Overall",
    *(SUBTASK_LABELS[name] for name in SUBTASKS),
    "Direction Grouped",
    "Total Samples",
    *(f"{SUBTASK_LABELS[name]} Samples" for name in SUBTASKS),
    "Evaluation Directory",
    "Results JSON",
    "Detail JSON",
)


def write_csv(path: Path, headers: Iterable[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(headers), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _populate_sheet(worksheet, headers: tuple[str, ...], rows: list[dict[str, Any]]) -> None:
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 30
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_index, header in enumerate(headers, start=1):
        values = [str(header)] + [str(row.get(header, "")) for row in rows]
        width = min(max(len(value) for value in values) + 2, 42)
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(width, 10)
        if header in {"1%", "5%", "25%", "50%", "100% / Final", "Overall", "Direction Grouped"} or header in SUBTASK_LABELS.values():
            for cell in worksheet.iter_cols(
                min_col=column_index,
                max_col=column_index,
                min_row=2,
                max_row=worksheet.max_row,
            ):
                for item in cell:
                    item.number_format = "0.0000"
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def write_workbook(
    path: Path,
    summaries: list[dict[str, Any]],
    details: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    _populate_sheet(summary_sheet, SUMMARY_HEADERS, summaries)
    detail_sheet = workbook.create_sheet("Subtask Detail")
    _populate_sheet(detail_sheet, DETAIL_HEADERS, details)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    args = parse_args()
    runs = discover_runs(args.eval_root.resolve(), args.reproduction_dir)
    detail_rows = [load_run(spec) for spec in runs]
    id_order = {architecture_id: index for index, architecture_id in enumerate(SUMMARY_IDS)}
    stage_order = {stage: index for index, stage in enumerate(CONTROLLED_STAGES)}
    detail_rows.sort(key=lambda row: (id_order[row["ID"]], stage_order[row["Stage Key"]]))
    summaries = summary_rows(detail_rows)

    prefix = args.output_prefix.resolve()
    summary_csv = prefix.with_name(f"{prefix.name}_summary.csv")
    detail_csv = prefix.with_name(f"{prefix.name}_subtask_detail.csv")
    workbook_path = prefix.with_suffix(".xlsx")
    write_csv(summary_csv, SUMMARY_HEADERS, summaries)
    write_csv(detail_csv, DETAIL_HEADERS, detail_rows)
    write_workbook(workbook_path, summaries, detail_rows)

    print(f"Validated evaluations: {len(detail_rows)}")
    print(f"Summary rows: {len(summaries)}")
    print(f"Summary CSV: {summary_csv}")
    print(f"Subtask CSV: {detail_csv}")
    print(f"Workbook: {workbook_path}")


if __name__ == "__main__":
    main()
